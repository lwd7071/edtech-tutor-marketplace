from pathlib import Path

from docx import Document
from docx.shared import Pt


root = Path(r"D:\EdTech")
path = root / "deliverables" / "Tutor_Match_Thuyet_Minh_Du_An_UTE_SV_STARTUP_2026_Ban_Trinh_Bay.docx"
doc = Document(path)


def replace_paragraph(paragraph, text):
    old_run = paragraph.runs[0] if paragraph.runs else None
    props = {
        "bold": old_run.bold if old_run else None,
        "italic": old_run.italic if old_run else None,
        "underline": old_run.underline if old_run else None,
        "font_name": old_run.font.name if old_run else "Times New Roman",
        "font_size": old_run.font.size if old_run else Pt(14),
        "font_color": old_run.font.color.rgb if old_run and old_run.font.color else None,
    }
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = props["bold"]
    run.italic = props["italic"]
    run.underline = props["underline"]
    run.font.name = props["font_name"]
    run.font.size = props["font_size"]
    if props["font_color"]:
        run.font.color.rgb = props["font_color"]


forecast = next(p for p in doc.paragraphs if p.text.startswith("Dự phóng 12 tháng:"))
replace_paragraph(
    forecast,
    forecast.text.replace("6 triệu đồng/tháng", "2 triệu đồng/tháng"),
)

feasibility = next(p for p in doc.paragraphs if p.text.startswith("Kịch bản tài chính bên dưới"))
replace_paragraph(feasibility, feasibility.text.replace("Chi phí 6 triệu đồng/tháng", "Chi phí 2 triệu đồng/tháng"))

funding = next(p for p in doc.paragraphs if p.text.startswith("Nhu cầu vốn theo kịch bản"))
replace_paragraph(funding, funding.text.replace("45 triệu đồng/12 tháng", "8,4 triệu đồng/12 tháng"))

break_even = next(p for p in doc.paragraphs if p.text.startswith("Điểm hòa vốn vận hành"))
replace_paragraph(
    break_even,
    "Điểm hòa vốn vận hành khoảng 34 gói hoàn tất tương đương/tháng: 2.000.000 / (70.000 hoa hồng bình quân − 10.000 chi phí biến đổi). Theo lộ trình dự phóng, số gói tháng vượt ngưỡng này từ tháng 4; dòng tiền lũy kế chuyển dương khoảng tháng 6. Kết quả dự phóng năm đầu dương 22,8 triệu đồng và mức thiếu hụt tiền mặt cao nhất khoảng 2,4 triệu đồng. Nếu duy trì 120 gói/tháng, phần đóng góp sau chi phí ước 5,2 triệu đồng/tháng. Nhu cầu vốn lưu động tham khảo khoảng 8,4 triệu đồng (2,4 triệu thiếu hụt cao nhất + 6 triệu dự phòng chi phí cố định ba tháng). Các khoản này chưa tính lương nhà sáng lập và chưa thay thế báo giá thực tế; đây không phải cam kết lợi nhuận.",
)

forecast_table = doc.tables[4]
updated_totals = ["−2.400.000", "+3.000.000", "+8.400.000", "+13.800.000", "+22.800.000"]
for row, total in zip(forecast_table.rows[1:], updated_totals):
    replace_paragraph(row.cells[3].paragraphs[0], total)

temp = path.with_name(path.stem + "_updated.docx")
doc.save(temp)
print(temp)
