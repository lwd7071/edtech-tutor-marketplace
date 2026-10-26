from pathlib import Path

from docx import Document
from docx.shared import Inches


source = Path(r"D:\EdTech\deliverables\Tutor_Match_Thuyet_Minh_Du_An_UTE_SV_STARTUP_2026_Ban_Trinh_Bay_updated.docx")
output = source.with_name("Tutor_Match_Thuyet_Minh_Du_An_UTE_SV_STARTUP_2026_Ban_Trinh_Bay_polished.docx")
doc = Document(source)

# Remove empty spacer paragraphs that pushed the cover date to its own page.
for paragraph in list(doc.paragraphs):
    xml = paragraph._p.xml
    if not paragraph.text.strip() and 'w:type="page"' not in xml and '<w:drawing' not in xml:
        paragraph._element.getparent().remove(paragraph._element)

# Keep the added student IDs on one line while retaining 14 pt text.
cover_roster = doc.tables[0]
cover_roster.autofit = False
cover_roster.alignment = 1  # centered
for column, width in zip(cover_roster.columns, (0.7, 2.5, 1.4)):
    column.width = Inches(width)
for row in cover_roster.rows:
    for cell, width in zip(row.cells, (0.7, 2.5, 1.4)):
        cell.width = Inches(width)

doc.save(output)
print(output)
